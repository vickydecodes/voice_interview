from __future__ import annotations
import io
import logging
from typing import Dict, Any
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import pypdf
except ImportError:
    pypdf = None

from docx import Document

from backend.models.user import ParsedResume
from backend.services.ai_service import ai_service, OpenRouterError

logger = logging.getLogger("resume_parser")


def extract_text_from_pdf(content: bytes) -> str:
    """Extracts raw text from a PDF file using PyMuPDF, pdfplumber, or pypdf."""
    # 1. Try PyMuPDF (fitz)
    if fitz:
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            text_parts = [page.get_text("text").strip() for page in doc if page.get_text("text").strip()]
            doc.close()
            full_text = "\n\n".join(text_parts).strip()
            if full_text:
                return full_text
        except Exception as e:
            logger.warning("PyMuPDF extraction failed: %s, trying fallback...", e)

    # 2. Try pdfplumber
    if pdfplumber:
        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                text_parts = [p.extract_text() for p in pdf.pages if p.extract_text()]
                full_text = "\n\n".join(text_parts).strip()
                if full_text:
                    return full_text
        except Exception as e:
            logger.warning("pdfplumber extraction failed: %s, trying fallback...", e)

    # 3. Try pypdf
    if pypdf:
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            text_parts = [page.extract_text() for page in reader.pages if page.extract_text()]
            full_text = "\n\n".join(text_parts).strip()
            if full_text:
                return full_text
        except Exception as e:
            logger.error("pypdf extraction failed: %s", e)

    raise ValueError("No readable text could be extracted from this PDF. Please ensure it is not an image-only scan.")


def extract_text_from_docx(content: bytes) -> str:
    """Extracts raw text from a DOCX file using python-docx."""
    try:
        file_stream = io.BytesIO(content)
        doc = Document(file_stream)
        text_parts = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
                
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
                    
        full_text = "\n".join(text_parts).strip()
        if not full_text:
            raise ValueError("No readable text could be extracted from this DOCX document.")
        return full_text
    except Exception as e:
        logger.error("DOCX extraction error: %s", e)
        raise ValueError(f"Failed to read DOCX document: {str(e)}")


RESUME_PARSER_SYSTEM_PROMPT = """You are an expert HR resume parser and talent analyst.
Extract structured profile information from the candidate's resume text.
Extract their name, email, phone number, skills list, education history, work experience (with role, company, duration, and bullet highlights), notable projects (with name, description, and tech stack), certifications, and a 2-3 sentence executive summary.

If a field is not present in the resume, provide an empty list or a sensible fallback (e.g. if email is missing, use 'candidate@simulator.local').
Be accurate, thorough, and do not invent experience not mentioned in the resume."""

SCHEMA_HINT = """{
  "name": "Candidate Full Name",
  "email": "candidate@example.com",
  "phone": "+1 555-0199",
  "skills": ["Python", "FastAPI", "React", "Docker"],
  "education": [
    {
      "degree": "B.S. in Computer Science",
      "institution": "Stanford University",
      "year": "2018 - 2022",
      "details": "Graduated Magna Cum Laude"
    }
  ],
  "experience": [
    {
      "role": "Senior Software Engineer",
      "company": "Tech Solutions",
      "duration": "2022 - Present",
      "highlights": [
        "Architected distributed microservices processing 1M daily requests",
        "Reduced latency by 45% using Redis caching"
      ]
    }
  ],
  "projects": [
    {
      "name": "Distributed Job Scheduler",
      "description": "High-throughput task queue in Go and Redis",
      "tech_stack": ["Go", "Redis", "gRPC"]
    }
  ],
  "certifications": [
    "AWS Certified Solutions Architect"
  ],
  "summary": "Results-driven engineer with 4+ years of experience building scalable backend architectures and distributed systems."
}"""


async def parse_resume_content(filename: str, content: bytes) -> ParsedResume:
    """Extracts text based on file format and uses AI to generate structured ParsedResume."""
    filename_lower = filename.lower()
    if filename_lower.endswith(".pdf"):
        raw_text = extract_text_from_pdf(content)
    elif filename_lower.endswith(".docx"):
        raw_text = extract_text_from_docx(content)
    else:
        raise ValueError("Unsupported file format. Please upload a PDF (.pdf) or Word document (.docx).")

    # Limit max raw text length if massive
    if len(raw_text) > 40000:
        raw_text = raw_text[:40000]

    user_prompt = f"Resume Content:\n\n{raw_text}"
    
    parsed_json = await ai_service.generate_json(
        system=RESUME_PARSER_SYSTEM_PROMPT,
        user=user_prompt,
        schema_hint=SCHEMA_HINT,
    )
    
    # Standardize email
    email = str(parsed_json.get("email") or "").strip().lower()
    if not email or "@" not in email:
        # Fallback if AI couldn't find an email
        safe_name = re.sub(r'[^a-zA-Z0-9]', '', str(parsed_json.get("name", "candidate")).lower())
        email = f"{safe_name or 'candidate'}@simulator.local"
    parsed_json["email"] = email

    # Parse into Pydantic model
    return ParsedResume(**parsed_json)
